"""Upload-store demo corpus — generate sample documents, encrypt, register them.

These six documents are the closed-network corpus served through the
uploads connector (still seeded by demo-seed/build_seed.py into
demo-seed/seed/). Company A content comes from the platform simulators —
see scripts/sync_demo.py.
"""

import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fileguard import db, parsers, security, store  # noqa: E402

# plaintext originals land here for dev/demo inspection (see folder README)
DEMO_DIR = Path(__file__).resolve().parent.parent / "dev_demo_files"

# name, kind, base_clearance, groups, builder()
SAMPLES: list[tuple[str, str, str, list[str], object]] = []


def _pdf(pages: list[str]) -> bytes:
    import pymupdf as fitz

    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx(paragraphs: list[str]) -> bytes:
    import docx

    d = docx.Document()
    for text in paragraphs:
        d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


SAMPLES = [
    ("employee_handbook.pdf", "pdf", "open", [], lambda: _pdf([
        "Employee Handbook\n"
        "Section 1: Working hours are 9am to 6pm, Monday to Friday.\n"
        "Section 2: All employees receive 20 days of annual leave.",
        "Section 3: Expense claims must be submitted within 30 days.\n"
        "Section 4: The canteen serves lunch from 11:30am to 2pm.",
    ])),
    ("q3_financials.xlsx", "xlsx", "confidential", ["finance"], lambda: _xlsx({
        "Summary": [
            ["Metric", "Q2", "Q3"],
            ["Revenue", "1,200,000", "1,450,000"],
            ["Costs", "900,000", "980,000"],
            ["Profit", "300,000", "470,000"],
        ],
        "Notes": [
            ["Note"],
            ["Q3 revenue boosted by the Meridian contract renewal."],
            ["Costs rose 8.9% due to cloud infrastructure expansion."],
        ],
    })),
    ("hr_policy.pdf", "pdf", "secret", ["HR"], lambda: _pdf([
        "HR Policy: Grievance Handling\n"
        "Grievances must be filed with the HR lead within 10 working days.\n"
        "The HR lead acknowledges within 2 working days.",
        "Escalation: unresolved grievances go to the arbitration panel.\n"
        "Retaliation against grievance filers is a dismissible offense.",
    ])),
    ("hr_incident_report.docx", "docx", "top_secret", ["HR"], lambda: _docx([
        "Incident Report IR-2041 — CONFIDENTIAL INVESTIGATION",
        "On 12 September an anonymous complaint was filed against a senior "
        "manager in Finance alleging misuse of expense accounts.",
        "Investigation status: evidence collected; hearing scheduled for "
        "October. Details restricted to HR leads.",
    ])),
    ("product_roadmap.docx", "docx", "secret", [], lambda: _docx([
        "Product Roadmap (Internal)",
        "Q4: launch the audit-trail viewer and reconciliation dashboard.",
        "Q1 next year: anomaly detection over access patterns.",
    ])),
    ("board_compensation.xlsx", "xlsx", "top_secret", ["finance"],
     lambda: _xlsx({
        "Remuneration": [
            ["Director", "Role", "FY2025 Total (SGD)"],
            ["Tan Wei Ming", "Chairman", "1,850,000"],
            ["Priya Nair", "CEO", "2,400,000"],
            ["Daniel Ong", "CFO", "1,450,000"],
            ["Sarah Lim", "Independent Director", "380,000"],
        ],
        "Notes": [
            ["Note"],
            ["CEO package includes a Meridian renewal performance bonus."],
            ["Figures are FY2025 audited totals incl. deferred awards."],
        ],
    })),
]


def main() -> None:
    conn = db.central()
    try:
        for name, kind, clearance, groups, builder in SAMPLES:
            plaintext = builder()
            assert parsers.detect_kind(name) == kind
            file_id = store.save_plaintext(plaintext)
            DEMO_DIR.mkdir(exist_ok=True)
            (DEMO_DIR / name).write_bytes(plaintext)  # dev-demo plaintext copy
            conn.execute(
                "INSERT OR REPLACE INTO files (file_id, name, kind, base_clearance, groups, label_date) "
                "VALUES (?,?,?,?,?,date('now'))",
                (file_id, name, kind, clearance, json.dumps(groups)),
            )
            spans = parsers.parse_to_spans(plaintext, kind)
            print(f"ingested {name:28} clearance={clearance:12} groups={groups} "
                  f"file_id={file_id[:12]}… spans={len(spans)}")
        conn.commit()
    finally:
        conn.close()


if __name__ == "__main__":
    main()
