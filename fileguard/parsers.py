"""Parsers: file bytes -> spans. Span = the unit the gatekeeper references.

  pdf   -> page spans        (start = end = 1-indexed page number)
  xlsx  -> cell-range spans  (start = "Sheet1!A1", end = "Sheet1!C10")
  docx  -> paragraph spans   (start = end = paragraph index, non-empty only)

The file itself is the single source of truth: extraction always re-parses
the plaintext, so the content-hash pin covers every span reference.
"""

import io

KIND_BY_SUFFIX = {".pdf": "pdf", ".xlsx": "xlsx", ".docx": "docx"}


class Span:
    __slots__ = ("kind", "start", "end", "text")

    def __init__(self, kind: str, start, end, text: str):
        self.kind, self.start, self.end, self.text = kind, start, end, text

    def matches(self, kind: str, start, end) -> bool:
        return self.kind == kind and self.start == start and self.end == end


def detect_kind(filename: str) -> str:
    for suffix, kind in KIND_BY_SUFFIX.items():
        if filename.lower().endswith(suffix):
            return kind
    raise ValueError(f"unsupported file type: {filename}")


def parse_to_spans(plaintext: bytes, kind: str) -> list[Span]:
    if kind == "pdf":
        return _parse_pdf(plaintext)
    if kind == "xlsx":
        return _parse_xlsx(plaintext)
    if kind == "docx":
        return _parse_docx(plaintext)
    raise ValueError(f"unknown kind {kind}")


def _parse_pdf(data: bytes) -> list[Span]:
    import pymupdf as fitz  # PyMuPDF

    spans = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                spans.append(Span("page", i, i, text))
    return spans


def _parse_xlsx(data: bytes) -> list[Span]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    spans = []
    for ws in wb.worksheets:
        rows = [row for row in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        # one span per non-empty contiguous row block, referenced A1-style
        block_start = None
        last_col = 0
        for r_idx, row in enumerate(rows, start=1):
            if any(c is not None and str(c).strip() for c in row):
                if block_start is None:
                    block_start = r_idx
                last_col = max(last_col, len(row))
            elif block_start is not None:
                spans.append(_xlsx_span(ws.title, block_start, r_idx - 1, last_col, rows))
                block_start = None
        if block_start is not None:
            spans.append(_xlsx_span(ws.title, block_start, len(rows), last_col, rows))
    wb.close()
    return spans


def _xlsx_span(sheet: str, r1: int, r2: int, n_cols: int, rows) -> Span:
    from openpyxl.utils import get_column_letter

    col = get_column_letter(n_cols)
    lines = []
    for r in range(r1, r2 + 1):
        cells = ["" if c is None else str(c) for c in rows[r - 1][:n_cols]]
        lines.append(" | ".join(cells).strip(" |"))
    text = "\n".join(line for line in lines if line)
    return Span("cell_range", f"{sheet}!A{r1}", f"{sheet}!{col}{r2}", text)


def _parse_docx(data: bytes) -> list[Span]:
    import docx

    document = docx.Document(io.BytesIO(data))
    spans = []
    for i, para in enumerate(document.paragraphs):
        text = para.text.strip()
        if text:
            spans.append(Span("paragraph", i, i, text))
    return spans


def extract_span(plaintext: bytes, kind: str, span_kind: str, start, end) -> str:
    """Copy the referenced span verbatim from the plaintext."""
    for span in parse_to_spans(plaintext, kind):
        if span.matches(span_kind, start, end):
            return span.text
    raise KeyError(f"span {span_kind}[{start}..{end}] not found in document")
